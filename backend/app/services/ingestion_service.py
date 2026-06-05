import os
import uuid
import re
import json
from dataclasses import dataclass
from pypdf import PdfReader
from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from qdrant_client.models import PointStruct
from sqlmodel import Session, select

from ..models.document import Document, DocumentChunk
from ..core.config import settings
from .embedding_service import generate_embeddings
from .ocr_service import OcrResult, convert_pdf_pages_to_images, ocr_image_file, ocr_pil_image_to_text
from .policy_citation_metadata import build_citation_metadata
from .qdrant_service import upload_points
from .source_risk_service import classify_source_risk


SECTION_RE = re.compile(
    r"\b(?:section|sec\.?|clause|article|chapter|part)\s+([0-9]+(?:\.[0-9]+)*)\b|"
    r"\b(?:दफा|परिच्छेद|बुँदा)\s*([०-९0-9]+(?:[.\-][०-९0-9]+)*)",
    re.IGNORECASE,
)
DEVANAGARI_RE = re.compile(r"[\u0900-\u097F]")
DEVANAGARI_MARK_AFTER_NON_DEVANAGARI_RE = re.compile(r"(^|[^\u0900-\u097F])[\u093A-\u094D]")
OCR_LATIN_NOISE_RE = re.compile(r"[A-Za-z%]")
LATIN_RE = re.compile(r"[A-Za-z]")
LEGACY_NEPALI_SYMBOL_RE = re.compile(r"[;\[\]\{\}\|ˆ÷]")
LEGACY_NEPALI_CLUSTER_RE = re.compile(
    r"(?:[;:][A-Za-z]|[A-Za-z][\]\[\{\}\|]|[A-Za-z]['ˆ÷]|[A-Za-z]/[A-Za-z]|[A-Za-z][+!@#$%^&*=][A-Za-z]?)"
)
LEGACY_NEPALI_TOKEN_RE = re.compile(
    r"(?:;f|;DaGw|sf|df|kg|k\||u/|g\]|n\]|O\{|lj|tyf|kq|u\||x'|z'|sfg'g|lgsfo|u/L|u\{|"
    r"ˆof|;DaGw|j\]|cf|clwsf/L|ljQLo|u\|fxs|k\|ltzt|k\|rlnt)",
    re.IGNORECASE,
)
COMMON_ENGLISH_WORD_RE = re.compile(
    r"\b(?:the|and|for|with|from|bank|customer|account|policy|section|shall|must|website|"
    r"verification|document|guarantee|published|handled|according|details)\b",
    re.IGNORECASE,
)
DEGRADED_NEPALI_TEXT_PATTERNS = (
    re.compile(r"वव[षशकदधचजतथन]"),
    re.compile(r"मम[ितधन]"),
    re.compile(r"[अ-ह] ु"),
    re.compile(r"\b(?:गन|अथ|काय|काम|पछ|समे|िेत्र|िोवक|भन्)\s"),
)


@dataclass(frozen=True)
class ChunkProfile:
    name: str
    chunk_size: int
    chunk_overlap: int


DEFAULT_CHUNK_PROFILE = ChunkProfile("default_text", 1000, 200)
REGULATORY_CHUNK_PROFILE = ChunkProfile("regulatory_section", 900, 180)
SPREADSHEET_CHUNK_PROFILE = ChunkProfile("spreadsheet_table", 1600, 120)
PRESENTATION_CHUNK_PROFILE = ChunkProfile("presentation_slide", 900, 120)
OCR_CHUNK_PROFILE = ChunkProfile("ocr_compact", 800, 100)

REGULATORY_DOCUMENT_TYPES = {
    "policy",
    "procedure",
    "manual",
    "compliance",
    "circular",
    "directive",
    "law",
    "act",
    "sop",
}


def _normalize_file_type(file_type: str | None) -> str:
    return (file_type or "").strip().lower().lstrip(".")


def _has_ocr_confidence(pages: list[dict] | None) -> bool:
    return any(page.get("ocr_confidence") is not None for page in pages or [])


def _has_dense_section_markers(pages: list[dict] | None) -> bool:
    text = "\n".join((page.get("text") or "")[:2500] for page in pages or [])
    return len(SECTION_RE.findall(text)) >= 2


def resolve_chunk_profile(
    *,
    file_type: str | None = None,
    document_type: str | None = None,
    pages: list[dict] | None = None,
    chunk_size: int | None = None,
    chunk_overlap: int | None = None,
) -> ChunkProfile:
    if chunk_size is not None or chunk_overlap is not None:
        size = chunk_size if chunk_size is not None else DEFAULT_CHUNK_PROFILE.chunk_size
        overlap = chunk_overlap if chunk_overlap is not None else min(DEFAULT_CHUNK_PROFILE.chunk_overlap, size // 5)
        return ChunkProfile("custom", size, min(overlap, max(size - 1, 0)))

    normalized_file_type = _normalize_file_type(file_type)
    normalized_document_type = (document_type or "").strip().lower()

    if normalized_file_type in {"xlsx", "xls", "csv"}:
        return SPREADSHEET_CHUNK_PROFILE
    if normalized_file_type in {"pptx", "ppt"}:
        return PRESENTATION_CHUNK_PROFILE
    if normalized_file_type in {"jpg", "jpeg", "png"} or _has_ocr_confidence(pages):
        return OCR_CHUNK_PROFILE
    if normalized_document_type in REGULATORY_DOCUMENT_TYPES or _has_dense_section_markers(pages):
        return REGULATORY_CHUNK_PROFILE
    return DEFAULT_CHUNK_PROFILE


def _extract_section_label(text: str | None) -> str | None:
    if not text:
        return None
    match = SECTION_RE.search(text[:1200])
    if not match:
        return None
    return " ".join(match.group(0).strip().split())[:80]


def _page_payload(
    *,
    page_number: int | None,
    text: str,
    extraction_confidence: float | None = None,
    ocr_confidence: float | None = None,
    table_confidence: float | None = None,
    page_bbox_json: str | None = None,
    section_label: str | None = None,
    table_metadata: dict | None = None,
    pdf_text_layer_repaired: bool = False,
    vision_transcription: bool = False,
    vision_model: str | None = None,
) -> dict:
    payload = {
        "page_number": page_number,
        "text": text,
        "extraction_confidence": extraction_confidence,
        "ocr_confidence": ocr_confidence,
        "table_confidence": table_confidence,
        "page_bbox_json": page_bbox_json,
        "section_label": section_label,
        "table_metadata": table_metadata,
    }
    if pdf_text_layer_repaired:
        payload["pdf_text_layer_repaired"] = True
    if vision_transcription:
        payload["vision_transcription"] = True
        payload["vision_model"] = vision_model or settings.LLM_C_MODEL
    return payload


def _degraded_devanagari_score(text: str) -> int:
    devanagari_count = len(DEVANAGARI_RE.findall(text or ""))
    if devanagari_count < 40:
        return 0

    score = 0
    score += len(DEVANAGARI_MARK_AFTER_NON_DEVANAGARI_RE.findall(text)) * 3
    score += len(re.findall(r"\s[\u093A-\u094D]", text)) * 2
    for pattern in DEGRADED_NEPALI_TEXT_PATTERNS:
        score += len(pattern.findall(text)) * 2
    return score


def _is_degraded_devanagari_pdf_text(text: str) -> bool:
    devanagari_count = len(DEVANAGARI_RE.findall(text or ""))
    if devanagari_count < 40:
        return False
    threshold = max(8, devanagari_count // 90)
    return _degraded_devanagari_score(text) >= threshold


def _legacy_nepali_text_layer_score(text: str) -> int:
    sample = (text or "")[:4000]
    if not sample.strip():
        return 0

    devanagari_count = len(DEVANAGARI_RE.findall(sample))
    if devanagari_count >= max(12, len(sample) // 50):
        return 0

    latin_count = len(LATIN_RE.findall(sample))
    if latin_count < 25:
        return 0

    legacy_symbol_count = len(LEGACY_NEPALI_SYMBOL_RE.findall(sample))
    cluster_count = len(LEGACY_NEPALI_CLUSTER_RE.findall(sample))
    legacy_token_count = len(LEGACY_NEPALI_TOKEN_RE.findall(sample))
    tokens = re.findall(r"[A-Za-z][A-Za-z'/{}\[\]|ˆ÷+]*", sample)
    odd_token_count = 0
    for token in tokens:
        if (
            re.search(r"[;{}\[\]|ˆ÷']", token)
            or LEGACY_NEPALI_TOKEN_RE.search(token)
            or (len(token) > 3 and not re.search(r"[aeiouAEIOU]", token))
        ):
            odd_token_count += 1

    score = legacy_symbol_count * 2 + cluster_count * 4 + legacy_token_count * 5
    if tokens:
        odd_ratio = odd_token_count / len(tokens)
        if odd_ratio >= 0.35:
            score += int(odd_ratio * 20)
        english_ratio = len(COMMON_ENGLISH_WORD_RE.findall(sample)) / len(tokens)
        if english_ratio >= 0.35 and legacy_symbol_count < 6 and legacy_token_count < 2:
            score -= 20
    if legacy_symbol_count / max(len(sample), 1) >= 0.035:
        score += 8
    if "ˆ" in sample or "÷" in sample:
        score += 10
    return max(score, 0)


def _is_legacy_nepali_pdf_text(text: str) -> bool:
    sample = (text or "")[:4000]
    latin_count = len(LATIN_RE.findall(sample))
    if latin_count < 25:
        return False
    threshold = max(18, min(80, latin_count // 35))
    return _legacy_nepali_text_layer_score(sample) >= threshold


def _is_usable_legacy_nepali_repair(text: str) -> bool:
    normalized = (text or "").strip()
    if not normalized:
        return False
    devanagari_count = len(DEVANAGARI_RE.findall(normalized))
    if devanagari_count < 12:
        return False
    if _is_legacy_nepali_pdf_text(normalized):
        return False
    latin_count = len(LATIN_RE.findall(normalized))
    return devanagari_count >= latin_count or devanagari_count >= 40


def _ocr_line_noise_score(line: str) -> int:
    return len(OCR_LATIN_NOISE_RE.findall(line or ""))


def _normalize_common_nepali_ocr_errors(text: str) -> str:
    normalized = re.sub(r"(?<=[\u0900-\u097F])\s+a\s+(?=[\u0900-\u097F])", " वा ", text)
    normalized = re.sub(r"(?<=[०-९])%+(?=[०-९])", "", normalized)
    return normalized


HANDWRITING_TRANSCRIPTION_PROMPT = """Transcribe the visible handwritten or printed text in this document image.

Return only the visible text line by line.
Preserve the original script and wording.
Do not translate, summarize, explain, or add labels.
If a word is unclear, write (unclear)."""


def _looks_like_failed_image_ocr(result: OcrResult) -> bool:
    text = (result.text or "").strip()
    if not text:
        return True
    confidence = result.confidence
    if confidence is not None and confidence >= settings.OCR_HANDWRITING_FALLBACK_CONFIDENCE:
        return False
    devanagari_count = len(DEVANAGARI_RE.findall(text))
    latin_count = len(LATIN_RE.findall(text))
    symbol_count = len(re.findall(r"[=|~><(){}\\[\\]_/\\\\]", text))
    if confidence is not None and confidence < settings.OCR_HANDWRITING_FALLBACK_CONFIDENCE:
        return True
    return latin_count > max(devanagari_count * 2, 12) and symbol_count >= 3


def _usable_vision_transcription(text: str | None) -> bool:
    if not text:
        return False
    normalized = text.strip()
    if not normalized:
        return False
    lowered = normalized.lower()
    failure_markers = (
        "failed to transcribe",
        "failed to analyze",
        "cannot transcribe",
        "can't transcribe",
        "unable to transcribe",
        "no visible text",
    )
    return not any(marker in lowered for marker in failure_markers)


def _transcribe_image_with_lipicore(file_path: str) -> str:
    from .llm_service import call_lipicore_vision_file

    return call_lipicore_vision_file(HANDWRITING_TRANSCRIPTION_PROMPT, file_path).strip()


def _maybe_replace_low_confidence_image_ocr(file_path: str, result: OcrResult) -> tuple[OcrResult, bool]:
    if not settings.OCR_HANDWRITING_FALLBACK_ENABLED:
        return result, False
    if not _looks_like_failed_image_ocr(result):
        return result, False

    transcription = _transcribe_image_with_lipicore(file_path)
    if not _usable_vision_transcription(transcription):
        return result, False
    return OcrResult(text=transcription, confidence=result.confidence), True


def _merge_direct_lines_for_ocr_noise(direct_text: str, ocr_text: str) -> str:
    direct_lines = direct_text.splitlines()
    ocr_lines = ocr_text.splitlines()
    merged_lines: list[str] = []
    for index, ocr_line in enumerate(ocr_lines):
        direct_line = direct_lines[index] if index < len(direct_lines) else ""
        if (
            direct_line.strip()
            and _ocr_line_noise_score(ocr_line) > 0
            and _ocr_line_noise_score(direct_line) == 0
            and len(DEVANAGARI_RE.findall(direct_line)) >= 3
            and _degraded_devanagari_score(direct_line) <= _degraded_devanagari_score(ocr_line) + 2
        ):
            merged_lines.append(direct_line)
        else:
            merged_lines.append(ocr_line)
    return "\n".join(merged_lines)


def _ocr_pdf_page_if_better(file_path: str, page_number: int, text: str) -> OcrResult | None:
    is_legacy_nepali_text_layer = _is_legacy_nepali_pdf_text(text)
    is_degraded_devanagari_text_layer = _is_degraded_devanagari_pdf_text(text)
    if not (is_legacy_nepali_text_layer or is_degraded_devanagari_text_layer):
        return None

    max_repair_pages = (
        settings.OCR_TEXT_LAYER_REPAIR_MAX_PAGES
        if is_legacy_nepali_text_layer
        else settings.OCR_MAX_PAGES
    )
    if page_number > max_repair_pages:
        return None

    images = convert_pdf_pages_to_images(file_path, first_page=page_number, last_page=page_number)
    for image in images:
        try:
            result = ocr_pil_image_to_text(image)
        finally:
            image.close()

        normalized_ocr_text = _normalize_common_nepali_ocr_errors(result.text)
        if is_legacy_nepali_text_layer and _is_usable_legacy_nepali_repair(normalized_ocr_text):
            return OcrResult(text=normalized_ocr_text, confidence=result.confidence)
        if (
            is_degraded_devanagari_text_layer
            and result.text.strip()
            and _degraded_devanagari_score(result.text) < _degraded_devanagari_score(text)
        ):
            return OcrResult(
                text=_merge_direct_lines_for_ocr_noise(text, normalized_ocr_text),
                confidence=result.confidence,
            )
    return None


def _extract_xlsx_pages(file_path: str) -> list[dict]:
    from openpyxl import load_workbook

    formula_wb = load_workbook(file_path, read_only=False, data_only=False)
    value_wb = load_workbook(file_path, read_only=False, data_only=True)
    pages = []
    try:
        for sheet in formula_wb.worksheets:
            value_sheet = value_wb[sheet.title]
            sheet_lines = [f"\n--- Sheet: {sheet.title} ---"]
            sheet_lines.append(f"Dimension: {sheet.calculate_dimension()}")
            merged_ranges = [str(cell_range) for cell_range in sheet.merged_cells.ranges]
            if sheet.merged_cells.ranges:
                sheet_lines.append(f"Merged ranges: {', '.join(merged_ranges)}")
            tables_metadata = []
            if sheet.tables:
                table_parts = []
                for name, table in sheet.tables.items():
                    ref = getattr(table, "ref", table)
                    tables_metadata.append({"name": name, "ref": str(ref)})
                    table_parts.append(f"{name}={ref}")
                tables = ", ".join(table_parts)
                sheet_lines.append(f"Tables: {tables}")
            formula_lines = []
            formula_cells = {}
            for row in sheet.iter_rows():
                cells = []
                for cell in row:
                    if cell.value is None:
                        continue
                    value = value_sheet[cell.coordinate].value
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formula_lines.append(f"{cell.coordinate} formula={cell.value} cached={value}")
                        formula_cells[cell.coordinate] = {
                            "formula": cell.value,
                            "cached": value,
                        }
                    cells.append(f"{cell.coordinate}={value if value is not None else cell.value}")
                if cells:
                    sheet_lines.append(" | ".join(cells))
            if formula_lines:
                sheet_lines.append("Formulas:")
                sheet_lines.extend(formula_lines)
            sheet_text = "\n".join(sheet_lines) + "\n"
            pages.append(_page_payload(
                page_number=None,
                text=sheet_text,
                extraction_confidence=0.98,
                table_confidence=0.9,
                section_label=f"Sheet: {sheet.title}",
                table_metadata={
                    "sheet_name": sheet.title,
                    "dimension": sheet.calculate_dimension(),
                    "merged_ranges": merged_ranges,
                    "tables": tables_metadata,
                    "formula_cells": formula_cells,
                },
            ))
        return pages
    finally:
        formula_wb.close()
        value_wb.close()


def extract_pages(file_path: str, file_type: str) -> list[dict]:
    pages = []
    text = ""
    ft = file_type.lower()
    if ft == 'pdf':
        pdf_page_count = 0
        try:
            import pdfplumber

            with pdfplumber.open(file_path) as pdf:
                pdf_page_count = len(pdf.pages)
                for index, page in enumerate(pdf.pages, start=1):
                    page_parts = []
                    page_text = page.extract_text() or ""
                    ocr_confidence = None
                    pdf_text_layer_repaired = False
                    if page_text.strip():
                        try:
                            ocr_result = _ocr_pdf_page_if_better(file_path, index, page_text)
                        except Exception:
                            ocr_result = None
                        if ocr_result is not None:
                            page_text = ocr_result.text
                            ocr_confidence = ocr_result.confidence
                            pdf_text_layer_repaired = True
                        page_parts.append(page_text)
                    if not pdf_text_layer_repaired:
                        try:
                            tables = page.extract_tables() or []
                        except Exception:
                            tables = []
                        for table_index, table in enumerate(tables, start=1):
                            page_parts.append(f"[Table {table_index} on page {index}]")
                            for row_index, row in enumerate(table or [], start=1):
                                values = [str(cell).strip() if cell is not None else "" for cell in row]
                                page_parts.append(f"Row {row_index}: " + " | ".join(values))
                    else:
                        tables = []
                    combined = "\n".join(part for part in page_parts if part.strip())
                    if combined.strip():
                        pages.append(_page_payload(
                            page_number=index,
                            text=combined,
                            extraction_confidence=0.82 if pdf_text_layer_repaired else 0.94 if page_text.strip() else 0.86,
                            ocr_confidence=ocr_confidence,
                            table_confidence=0.84 if tables else None,
                            pdf_text_layer_repaired=pdf_text_layer_repaired,
                        ))
                        text += combined + "\n"
        except Exception:
            pass

        try:
            reader = PdfReader(file_path)
            pdf_page_count = pdf_page_count or len(reader.pages)
            if not text.strip():
                for index, page in enumerate(reader.pages, start=1):
                    page_text = page.extract_text()
                    if page_text:
                        pages.append(_page_payload(
                            page_number=index,
                            text=page_text,
                            extraction_confidence=0.78,
                        ))
                        text += page_text + "\n"
        except Exception:
            pass  # PDF reading failed, will be caught by empty text check

        # If PDF has no extractable text, use open-source OCR on rendered pages.
        if not text.strip() and pdf_page_count:
            try:
                last_page = min(settings.OCR_MAX_PAGES, pdf_page_count)
                images = convert_pdf_pages_to_images(file_path, first_page=1, last_page=last_page)
                for index, img in enumerate(images, start=1):
                    result = ocr_pil_image_to_text(img)
                    if not result.text.strip():
                        continue
                    pages.append(_page_payload(
                        page_number=index,
                        text=result.text,
                        extraction_confidence=0.82,
                        ocr_confidence=result.confidence,
                    ))
                    text += result.text + "\n"
            except Exception as exc:
                raise RuntimeError(f"Open-source OCR failed for scanned PDF: {exc}") from exc
    elif ft == 'docx':
        doc = DocxDocument(file_path)
        for para in doc.paragraphs:
            text += para.text + "\n"
        pages.append(_page_payload(page_number=None, text=text, extraction_confidence=1.0))
    elif ft in ['txt', 'csv']:
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
        pages.append(_page_payload(page_number=None, text=text, extraction_confidence=1.0))
    elif ft in ['xlsx', 'xls']:
        pages.extend(_extract_xlsx_pages(file_path))
        text += "\n".join(page["text"] for page in pages)
    elif ft in ['pptx', 'ppt']:
        from pptx import Presentation
        prs = Presentation(file_path)
        for i, slide in enumerate(prs.slides):
            slide_text = f"\n--- Slide {i+1} ---\n"
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    slide_text += shape.text + "\n"
            pages.append(_page_payload(
                page_number=i + 1,
                text=slide_text,
                extraction_confidence=0.95,
                section_label=f"Slide {i + 1}",
            ))
            text += slide_text
    elif ft in ['jpg', 'jpeg', 'png']:
        result = ocr_image_file(file_path)
        result, vision_transcription = _maybe_replace_low_confidence_image_ocr(file_path, result)
        text = result.text
        pages.append(_page_payload(
            page_number=1,
            text=text,
            extraction_confidence=0.68 if vision_transcription else 0.82,
            ocr_confidence=result.confidence,
            vision_transcription=vision_transcription,
            vision_model="LipiCore" if vision_transcription else None,
        ))
    if not pages and text.strip():
        pages.append(_page_payload(page_number=None, text=text, extraction_confidence=0.7))
    return pages


def extract_text(file_path: str, file_type: str) -> str:
    return "\n".join(page["text"] for page in extract_pages(file_path, file_type))


def build_indexable_chunks(
    pages: list[dict],
    chunk_size: int | None = None,
    chunk_overlap: int | None = None,
    file_type: str | None = None,
    document_type: str | None = None,
) -> list[dict]:
    profile = resolve_chunk_profile(
        file_type=file_type,
        document_type=document_type,
        pages=pages,
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=profile.chunk_size,
        chunk_overlap=profile.chunk_overlap,
        length_function=len,
    )
    indexable_chunks = []
    for page in pages:
        page_text = (page.get("text") or "").strip()
        if not page_text:
            continue
        for chunk_text in text_splitter.split_text(page_text):
            risk = classify_source_risk(chunk_text)
            citation = build_citation_metadata(
                page=page,
                chunk_text=chunk_text,
                document_type=document_type,
            )
            indexable_chunks.append({
                "text": chunk_text,
                "page_number": page.get("page_number"),
                "pdf_page_number": citation["pdf_page_number"],
                "printed_page_number": citation["printed_page_number"],
                "document_heading": citation["document_heading"],
                "clause_number": citation["clause_number"],
                "citation_confidence": citation["citation_confidence"],
                "citation_incomplete_reasons": citation["citation_incomplete_reasons"],
                "legal_hierarchy": citation.get("legal_hierarchy"),
                "section_label": page.get("section_label") or _extract_section_label(chunk_text),
                "extraction_confidence": page.get("extraction_confidence"),
                "ocr_confidence": page.get("ocr_confidence"),
                "table_confidence": page.get("table_confidence"),
                "page_bbox_json": page.get("page_bbox_json"),
                "table_metadata": page.get("table_metadata"),
                "source_risk_level": risk["risk_level"],
                "source_risk_flags": risk["flags"],
            })
    return indexable_chunks


def create_extraction_review_records(
    db: Session,
    *,
    bank_id: int,
    document_id: int,
    pages: list[dict],
) -> list:
    from ..models.document_intelligence import DocumentExtractionPage
    from .document_intelligence_service import build_extraction_flags
    from .extraction_quality_service import build_extraction_quality

    records = []
    for page in pages:
        quality = build_extraction_quality(
            extraction_confidence=page.get("extraction_confidence"),
            ocr_confidence=page.get("ocr_confidence"),
            table_confidence=page.get("table_confidence"),
            text=page.get("text"),
        )
        visual_flags = build_extraction_flags(
            ocr_confidence=page.get("ocr_confidence"),
            table_confidence=page.get("table_confidence"),
            layout_confidence=page.get("layout_confidence"),
            has_handwriting=bool(page.get("has_handwriting")),
            has_signature_like_region=bool(page.get("has_signature_like_region")),
            has_stamp_like_region=bool(page.get("has_stamp_like_region")),
        )
        flags = list(dict.fromkeys([*quality["flags"], *visual_flags]))
        requires_review = quality["quality_bucket"] == "review_required" or bool(visual_flags)
        record = DocumentExtractionPage(
            bank_id=bank_id,
            document_id=document_id,
            page_number=page.get("page_number") or 0,
            extracted_text=page.get("text") or "",
            parser=page.get("parser") or "ingestion",
            ocr_confidence=page.get("ocr_confidence"),
            table_confidence=page.get("table_confidence"),
            layout_confidence=page.get("layout_confidence"),
            page_image_path=page.get("page_image_path"),
            bbox_json=page.get("page_bbox_json"),
            flags_json=json.dumps(flags),
            review_status="pending" if requires_review else "verified",
        )
        db.add(record)
        records.append(record)
    db.commit()
    for record in records:
        db.refresh(record)
    return records


def auto_catalog(text: str, filename: str) -> dict:
    """
    Use LLM to automatically classify document type, department, and suggest title.
    Returns dict with keys: document_type, department, title.
    Falls back gracefully if LLM can't parse response.
    """
    import json
    from .llm_service import call_llm

    try:
        prompt = f"""Analyze this document excerpt and classify it. Respond ONLY as valid JSON with no other text.

Extract and return:
1. "document_type": exactly one of [policy, procedure, manual, compliance, report, circular, directive, other]
2. "department": one of [Credit, Treasury, Operations, HR, Compliance, Audit, IT, Risk, General]
3. "title": a concise document title (max 10 words)

Document filename: {filename}

First 3000 characters of content:
{text[:3000]}

Respond ONLY as JSON, no markdown, no explanation."""

        response = call_llm(prompt).strip()

        # Try to extract JSON from response (handles models that may add markdown formatting)
        if "```json" in response:
            response = response.split("```json")[1].split("```")[0].strip()
        elif "```" in response:
            response = response.split("```")[1].split("```")[0].strip()

        data = json.loads(response)

        return {
            "document_type": data.get("document_type", "other"),
            "department": data.get("department", "General"),
            "title": data.get("title", filename)
        }
    except Exception as e:
        # Gracefully fall back to defaults if LLM fails
        return {
            "document_type": "other",
            "department": "General",
            "title": filename
        }


def process_document(document_id: int):
    """Background task — creates its own DB session so it is safe to run after the request ends."""
    from ..db.session import engine
    from ..models.bank import Bank  # noqa: F401
    from ..models.chat import ChatSession  # noqa: F401
    from ..models.user import User  # noqa: F401

    with Session(engine) as db:
        try:
            doc = db.get(Document, document_id)
            if not doc:
                return

            doc.status = "extracting_text"
            doc.processing_progress = 10
            doc.processing_message = "Reading document text..."
            db.add(doc)
            db.commit()

            # 1. Extract text
            pages = extract_pages(doc.file_path, doc.file_type)
            text = "\n".join(page["text"] for page in pages)
            try:
                create_extraction_review_records(db, bank_id=doc.bank_id, document_id=doc.id, pages=pages)
            except Exception:
                pass

            doc.status = "chunking"
            doc.processing_progress = 30
            doc.processing_message = "Preparing document for search..."
            db.add(doc)
            db.commit()

            # 2. Split into page-aware chunks
            chunks = build_indexable_chunks(
                pages,
                file_type=doc.file_type,
                document_type=doc.document_type,
            )

            if not chunks:
                doc.status = "failed"
                doc.processing_message = "Failed to extract text."
                db.add(doc)
                db.commit()
                return

            doc.status = "embedding"
            doc.processing_progress = 60
            doc.processing_message = "Creating secure document index..."
            db.add(doc)
            db.commit()

            # 3. Generate embeddings in batches
            BATCH = 64
            embeddings = []
            for i in range(0, len(chunks), BATCH):
                embeddings.extend(generate_embeddings([chunk["text"] for chunk in chunks[i:i + BATCH]]))

            # 4. Store in Qdrant and DB (bulk insert)
            points = []
            chunk_records = []
            for i, (chunk, embedding) in enumerate(zip(chunks, embeddings)):
                point_id = str(uuid.uuid4())
                chunk_records.append(DocumentChunk(
                    bank_id=doc.bank_id,
                    document_id=doc.id,
                    chunk_index=i,
                    chunk_text=chunk["text"],
                    page_number=chunk["page_number"],
                    printed_page_number=chunk.get("printed_page_number"),
                    document_heading=chunk.get("document_heading"),
                    clause_number=chunk.get("clause_number"),
                    citation_confidence=chunk.get("citation_confidence"),
                    citation_incomplete_reasons_json=json.dumps(chunk.get("citation_incomplete_reasons", [])),
                    extraction_confidence=chunk.get("extraction_confidence"),
                    ocr_confidence=chunk.get("ocr_confidence"),
                    table_confidence=chunk.get("table_confidence"),
                    page_bbox_json=chunk.get("page_bbox_json"),
                    source_risk_level=chunk.get("source_risk_level", "low"),
                    source_risk_flags_json=json.dumps(chunk.get("source_risk_flags", [])),
                    department=doc.department,
                    access_level=doc.access_level or 0,
                    document_scope=doc.document_scope,
                    session_id=doc.session_id,
                    document_status=doc.status,
                    version_state=doc.version_state,
                    qdrant_point_id=point_id,
                ))
                points.append(PointStruct(
                    id=point_id,
                    vector=embedding,
                    payload={
                        "bank_id": doc.bank_id,
                        "document_id": doc.id,
                        "chunk_index": i,
                        "text": chunk["text"],
                        "page_number": chunk["page_number"],
                        "pdf_page_number": chunk.get("pdf_page_number"),
                        "printed_page_number": chunk.get("printed_page_number"),
                        "document_heading": chunk.get("document_heading"),
                        "clause_number": chunk.get("clause_number"),
                        "citation_confidence": chunk.get("citation_confidence"),
                        "citation_incomplete_reasons": chunk.get("citation_incomplete_reasons", []),
                        "legal_hierarchy": chunk.get("legal_hierarchy"),
                        "section_label": chunk["section_label"],
                        "extraction_confidence": chunk.get("extraction_confidence"),
                        "ocr_confidence": chunk.get("ocr_confidence"),
                        "table_confidence": chunk.get("table_confidence"),
                        "page_bbox_json": chunk.get("page_bbox_json"),
                        "source_risk_level": chunk.get("source_risk_level", "low"),
                        "source_risk_flags": chunk.get("source_risk_flags", []),
                        "department": doc.department,
                        "access_level": doc.access_level or 0,
                        "document_status": doc.status,
                        "version_state": doc.version_state,
                        "document_scope": doc.document_scope,
                        "session_id": doc.session_id,
                        "embedding_model": settings.EMBEDDING_MODEL,
                        "embedding_dimension": settings.EMBEDDING_DIMENSION,
                    }
                ))

            db.add_all(chunk_records)
            db.commit()

            doc.status = "indexing"
            doc.processing_progress = 90
            doc.processing_message = "Adding document to session..."
            db.add(doc)
            db.commit()

            # Upload to Qdrant in batches
            QDRANT_BATCH = 100
            for i in range(0, len(points), QDRANT_BATCH):
                upload_points(points[i:i + QDRANT_BATCH])

            # 5. Generate summary
            try:
                from .llm_service import call_llm
                summary_prompt = f"Summarize the following document concisely in 2-3 sentences:\n\n{text[:4000]}"
                doc.summary = call_llm(summary_prompt)
            except Exception:
                doc.summary = "Summary generation failed."

            doc.status = "ready"
            doc.processing_progress = 100
            doc.processing_message = "Document is ready. You can ask follow-up questions."
            db.add(doc)
            db.commit()

            for chunk_record in db.exec(select(DocumentChunk).where(DocumentChunk.document_id == doc.id)).all():
                chunk_record.department = doc.department
                chunk_record.access_level = doc.access_level or 0
                chunk_record.document_scope = doc.document_scope
                chunk_record.session_id = doc.session_id
                chunk_record.document_status = doc.status
                chunk_record.version_state = doc.version_state
                db.add(chunk_record)
            db.commit()
            try:
                from .qdrant_service import update_points_by_document_payload
                update_points_by_document_payload(doc.id, doc.bank_id, {
                    "department": doc.department,
                    "access_level": doc.access_level or 0,
                    "document_status": doc.status,
                    "version_state": doc.version_state,
                    "document_scope": doc.document_scope,
                    "session_id": doc.session_id,
                })
            except Exception:
                pass

            # Auto-catalog in background (non-blocking)
            try:
                from .llm_service import call_llm
                catalog = auto_catalog(text, doc.file_name)
                if doc.document_type == "other":
                    doc.document_type = catalog["document_type"]
                if not doc.department:
                    doc.department = catalog["department"]
                db.add(doc)
                db.commit()
                for chunk_record in db.exec(select(DocumentChunk).where(DocumentChunk.document_id == doc.id)).all():
                    chunk_record.department = doc.department
                    chunk_record.access_level = doc.access_level or 0
                    db.add(chunk_record)
                db.commit()
                try:
                    from .qdrant_service import update_points_by_document_payload
                    update_points_by_document_payload(doc.id, doc.bank_id, {
                        "department": doc.department,
                        "access_level": doc.access_level or 0,
                    })
                except Exception:
                    pass
            except Exception:
                pass  # Cataloging failure doesn't block document readiness

        except Exception as e:
            try:
                doc = db.get(Document, document_id)
                if doc:
                    doc.status = "failed"
                    doc.processing_message = f"Error: {str(e)}"
                    db.add(doc)
                    db.commit()
            except Exception:
                pass
            raise

import os
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Request
from fastapi.responses import JSONResponse
from sqlmodel import Session
from typing import Optional

from ..db.session import get_session
from ..api.deps import get_current_user
from ..models.user import User
from ..services.task_service import get_all_templates, get_template
from ..services.llm_service import async_call_llm_a, async_call_llm_b
from ..services.guardrail_service import detect_prompt_injection, detect_and_mask_pii
from ..services.audit_service import log_audit_event
from ..core.limiter import limiter

router = APIRouter()

ALLOWED_EXTENSIONS = {'.pdf', '.docx', '.txt', '.xlsx', '.xls', '.csv', '.pptx', '.ppt'}
MAX_UPLOAD_BYTES = 20 * 1024 * 1024  # 20 MB for task files
MAX_TASK_INPUT_CHARS = 12_000
TASK_HEAD_CHARS = 8_000
TASK_TAIL_CHARS = 3_000


def _fit_task_context(text: str) -> tuple[str, bool]:
    if len(text) <= MAX_TASK_INPUT_CHARS:
        return text, False

    notice = (
        f"[Content shortened for model context: kept the first {TASK_HEAD_CHARS} "
        f"and last {TASK_TAIL_CHARS} characters from {len(text)} total characters.]"
    )
    shortened = f"{text[:TASK_HEAD_CHARS]}\n\n{notice}\n\n{text[-TASK_TAIL_CHARS:]}"
    return shortened, True


def _read_file_content(file_path: str, filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    try:
        if ext == '.txt' or ext == '.csv':
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        elif ext == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            return "\n".join(p.extract_text() or "" for p in reader.pages)
        elif ext in ('.docx',):
            from docx import Document
            doc = Document(file_path)
            return "\n".join(p.text for p in doc.paragraphs)
        elif ext in ('.xlsx', '.xls'):
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows(values_only=True):
                    lines.append("\t".join(str(c) if c is not None else "" for c in row))
            return "\n".join(lines)
        else:
            return ""
    except Exception as e:
        return f"[Could not read file: {e}]"


@router.get("/templates")
def list_templates(current_user: User = Depends(get_current_user)):
    return get_all_templates()


@router.post("/run")
@limiter.limit("20/minute")
async def run_task(
    request: Request,
    template_id: str = Form(...),
    content: str = Form(default=""),
    extra: str = Form(default=""),
    file: Optional[UploadFile] = File(default=None),
    db: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
):
    template = get_template(template_id)
    if not template:
        raise HTTPException(status_code=404, detail="Task template not found")

    file_text = ""
    tmp_path = None
    if file and file.filename:
        ext = os.path.splitext(file.filename)[1].lower()
        if ext not in ALLOWED_EXTENSIONS:
            raise HTTPException(status_code=400, detail=f"File type '{ext}' not supported")

        import tempfile, uuid
        tmp_path = os.path.join(tempfile.gettempdir(), f"task_{uuid.uuid4().hex}{ext}")
        raw = await file.read()
        if len(raw) > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=400, detail="File exceeds 20 MB limit")
        with open(tmp_path, "wb") as f_out:
            f_out.write(raw)
        file_text = _read_file_content(tmp_path, file.filename)
        os.remove(tmp_path)

    # Combine file text + typed content
    combined_content = "\n\n".join(filter(None, [file_text, content])).strip()
    if not combined_content:
        raise HTTPException(status_code=400, detail="Provide text content or upload a file")

    # Guardrails
    full_input = f"{combined_content}\n{extra}"
    if detect_prompt_injection(full_input, db, current_user.id, current_user.bank_id):
        raise HTTPException(status_code=400, detail="Input contains disallowed content")

    masked_content = detect_and_mask_pii(combined_content)
    masked_extra = detect_and_mask_pii(extra) if extra else ""
    masked_content, input_shortened = _fit_task_context(masked_content)

    user_prompt = template.user_prompt_template.format(
        content=masked_content,
        extra=masked_extra or "N/A",
    )

    call_fn = async_call_llm_b if template.llm == "b" else async_call_llm_a
    result = await call_fn(
        user_prompt,
        system=template.system_prompt,
        user_id=current_user.id,
        role=current_user.role,
    )

    log_audit_event(
        db=db,
        action=f"task_{template_id}",
        resource_type="task",
        bank_id=current_user.bank_id,
        user_id=current_user.id,
        metadata={
            "template": template_id,
            "has_file": file is not None and file.filename != "",
            "input_shortened": input_shortened,
            "input_chars": len(combined_content),
        },
    )

    return {
        "template_id": template_id,
        "template_name": template.name,
        "result": result,
        "suggested_formats": template.suggested_formats,
    }
